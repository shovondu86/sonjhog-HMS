import io

import openpyxl
import pytest

from conftest import future, wd
from app.scheduling import today

XLS = ".xls"


def make_xlsx(rows_spec):
    """Mimics the OPD room-allocation sheet: header row, merged section headers, merged day cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["OPD Room Allocation"])
    ws.append([])
    ws.append(["Room", "Doctor", "Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
    r = 4
    for spec in rows_spec:
        kind = spec[0]
        if kind == "section":
            ws.cell(r, 1, spec[1])
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        elif kind == "row":
            _, room, name, days = spec  # days: list of (first_col_idx, last_col_idx, text)
            ws.cell(r, 1, room)
            ws.cell(r, 2, name)
            for c1, c2, text in days:
                ws.cell(r, c1, text)
                if c2 > c1:
                    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


GOOD = [
    ("section", "OPD - (3rd Floor) Desk - A"),
    # Saturday..Thursday merged (cols 3..8) = SIX sessions, Friday x
    ("row", "3042", "Prof. Dr. Khwaja Nazim Uddin  01711-223344", [(3, 8, "09:00 am To 05:00 pm"), (9, 9, "x")]),
    # two sessions per day, Sat..Wed merged
    ("row", "3043", "Dr. Two", [(3, 7, "09:00 am To 12:00 pm & 03:00 pm To 08:00 pm"), (8, 8, "X"), (9, 9, "-")]),
    ("row", "PFT Room", "", [(3, 9, "x")]),
    ("section", "OPD - (4th Floor) Desk - B"),
    ("row", "4001", "Dr. Three", [(3, 3, "10:00 am To 01:00 pm"), (4, 9, "x")]),
]


def upload(client, headers, data, name, mode, **params):
    return client.post(
        "/schedules/import", params={"mode": mode, **params}, headers=headers,
        files={"file": (name, data, "application/octet-stream")},
    )


@pytest.fixture()
def three_doctors(hosp):
    return [hosp.doctor("Prof. Dr. Khwaja Nazim Uddin"), hosp.doctor("Dr. Two"), hosp.doctor("Dr. Three")]


def test_xlsx_preview_then_apply(client, hosp, admin, three_doctors):
    khwaja, two, three = three_doctors
    data = make_xlsx(GOOD)

    # no capacity anywhere -> MISSING_CAPACITY errors, apply refused
    pre = upload(client, admin, data, "opd.xlsx", "preview").json()
    assert pre["summary"]["errors"] > 0 and {p["code"] for p in pre["problems"] if p["severity"] == "error"} == {"MISSING_CAPACITY"}
    refused = upload(client, admin, data, "opd.xlsx", "apply")
    assert refused.status_code == 409 and refused.json()["code"] == "IMPORT_HAS_ERRORS"
    assert refused.json()["candidates"]

    pre = upload(client, admin, data, "opd.xlsx", "preview", default_max_patients=20).json()
    assert pre["summary"]["errors"] == 0 and pre["applied"] is False
    assert pre["summary"]["doctors_in_file"] == 3 and pre["summary"]["doctors_changed"] == 3
    assert pre["summary"]["mobiles_removed"] == 1
    codes = {p["code"] for p in pre["problems"]}
    assert {"MOBILE_REMOVED", "FACILITY_ROW"} <= codes
    by = {d["doctor_name"]: d["sessions"] for d in pre["doctors"]}
    assert len(by["Prof. Dr. Khwaja Nazim Uddin"]) == 6  # merged Sat..Thu = six sessions
    assert len(by["Dr. Two"]) == 10  # two sessions x five days
    k = by["Prof. Dr. Khwaja Nazim Uddin"][0]
    assert (k["room"], k["floor"], k["desk"], k["max_patients"]) == ("3042", "3rd", "A", 20)
    assert by["Dr. Three"][0]["floor"] == "4th" and by["Dr. Three"][0]["desk"] == "B"
    # preview changed nothing
    assert client.get(f"/doctors/{khwaja['id']}/schedules", headers=admin).json() == []

    ok = upload(client, admin, data, "opd.xlsx", "apply", default_max_patients=20).json()
    assert ok["applied"] is True
    rows = client.get(f"/doctors/{khwaja['id']}/schedules", headers=admin).json()
    assert len(rows) == 6 and all(r["valid_from"] == str(today()) for r in rows)
    # audit: one import entry per doctor
    log = client.get("/audit-log", params={"entity_type": "doctor"}, headers=admin).json()
    assert sorted(e["action"] for e in log if e["action"] == "import") == ["import"] * 3

    # re-uploading the same file is a no-op
    again = upload(client, admin, data, "opd.xlsx", "preview").json()
    assert again["summary"]["doctors_changed"] == 0 and all(d["changed"] is False for d in again["doctors"])


def test_xlsx_errors(client, hosp, admin, three_doctors):
    bad = [
        ("section", "OPD - (3rd Floor) Desk - A"),
        ("row", "1", "Dr. Two", [(3, 3, ""), (4, 4, "9 to 5"), (5, 9, "x")]),       # blank + unparseable
        ("row", "2", "Dr. Nobody", [(3, 9, "09:00 am To 12:00 pm")]),               # unknown doctor
        ("row", "9", "Dr. Three", [(3, 3, "09:00 am To 12:00 pm"), (4, 9, "x")]),
        ("row", "9", "Prof. Dr. Khwaja Nazim Uddin", [(3, 3, "10:00 am To 11:00 am"), (4, 9, "x")]),  # same room, overlapping
    ]
    pre = upload(client, admin, make_xlsx(bad), "opd.xlsx", "preview", default_max_patients=10).json()
    got = {p["code"] for p in pre["problems"]}
    assert {"BLANK_HOURS", "UNPARSEABLE_HOURS", "UNKNOWN_DOCTOR", "ROOM_CONFLICT"} <= got


def test_ambiguous_and_not_in_file(client, hosp, admin):
    a, b = hosp.doctor("Dr. Same Name"), hosp.doctor("Dr. Same Name")
    solo = hosp.doctor("Dr. Solo")
    d = future("sunday")
    hosp.schedule(solo["id"], d)
    sheet = [("row", "1", "Dr. Same Name", [(3, 3, "09:00 am To 12:00 pm"), (4, 9, "x")])]
    pre = upload(client, admin, make_xlsx(sheet), "a.xlsx", "preview", default_max_patients=5).json()
    assert "AMBIGUOUS_DOCTOR" in {p["code"] for p in pre["problems"]}
    assert "NOT_IN_FILE" in {p["code"] for p in pre["problems"]}


def test_xls_format(client, hosp, admin, three_doctors):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("OPD")
    for i, h in enumerate(["Room", "Doctor", "Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]):
        ws.write(0, i, h)
    ws.write_merge(1, 1, 0, 8, "OPD - (2nd Floor) Desk - C")
    ws.write(2, 0, "2001")
    ws.write(2, 1, "Dr. Two")
    ws.write_merge(2, 2, 2, 4, "08:00 am To 02:00 pm")
    ws.write_merge(2, 2, 5, 8, "x")
    buf = io.BytesIO()
    wb.save(buf)
    pre = upload(client, admin, buf.getvalue(), "opd.xls", "preview", default_max_patients=8).json()
    s = pre["doctors"][0]["sessions"]
    assert len(s) == 3 and {x["weekday"] for x in s} == {"saturday", "sunday", "monday"}
    assert s[0]["floor"] == "2nd" and s[0]["desk"] == "C"


def test_csv_and_capacity_precedence(client, hosp, admin):
    doc = hosp.doctor("Dr. Csv")
    d = future("sunday")
    hosp.schedule(doc["id"], d, "09:00", "12:00", cap=17, slot_minutes=10, report_capacity=3)  # current capacity 17
    csv_text = (
        "doctor_id,weekday,start_time,end_time,room,floor,desk,max_patients,slot_minutes,report_capacity\n"
        f"{doc['id']},sunday,10:00,13:00,R1,1st,A,,,\n"           # no capacity in file -> current (17)
        f"{doc['id']},monday,10:00,13:00,R1,1st,A,25,20,2\n"      # file value wins
        f"{doc['id']},tuesday,10:00,13:00,R1,1st,A,,,\n"          # no current for tuesday -> default
    )
    pre = upload(client, admin, csv_text.encode(), "s.csv", "preview", default_max_patients=9).json()
    caps = {s["weekday"]: (s["max_patients"], s["slot_minutes"], s["report_capacity"]) for s in pre["doctors"][0]["sessions"]}
    assert caps["sunday"] == (17, 10, 3)   # carried over from the current Sunday row
    assert caps["monday"] == (25, 20, 2)
    assert caps["tuesday"][0] == 9
    assert pre["doctors"][0]["changed"] is True

    past = upload(client, admin, csv_text.encode(), "s.csv", "preview", valid_from=str(today().replace(year=today().year - 1)))
    assert past.status_code == 422
    assert upload(client, admin, b"x", "notes.txt", "preview").json()["code"] == "UNSUPPORTED_FILE"
    assert upload(client, admin, b"x", "notes.txt", "preview").status_code == 415


def test_apply_keeps_future_bookings_attached(client, hosp, admin):
    doc = hosp.doctor("Dr. Keep")
    d = future("sunday")
    hosp.schedule(doc["id"], d, "09:00", "12:00", cap=5, room="1", floor="1st", desk="A")
    p = hosp.patient()
    a = hosp.book(p["id"], doc["id"], d).json()

    # room changes, hours the same: the booking is re-linked, not "affected"
    csv_text = (
        "doctor_id,weekday,start_time,end_time,room,floor,desk,max_patients\n"
        f"{doc['id']},{wd(d)},09:00,12:00,2,1st,A,5\n"
    )
    res = upload(client, admin, csv_text.encode(), "s.csv", "apply").json()
    assert res["applied"] is True and res["affected_appointments"] == []
    new = [s for s in client.get(f"/doctors/{doc['id']}/schedules", headers=admin).json() if s["room"] == "2"][0]
    now = client.get(f"/appointments/{a['id']}", headers=admin).json()
    assert now["schedule_id"] == new["id"] and now["serial_number"] == 1
    av = client.get(f"/doctors/{doc['id']}/availability", params={"from": str(d), "to": str(d)}, headers=admin).json()["days"][0]
    assert av["sessions"][0]["booked"] == 1 and av["sessions"][0]["room"] == "2"

    # hours move away: the booking is reported and left alone
    csv_text = (
        "doctor_id,weekday,start_time,end_time,room,floor,desk,max_patients\n"
        f"{doc['id']},{wd(d)},15:00,18:00,2,1st,A,5\n"
    )
    res = upload(client, admin, csv_text.encode(), "s.csv", "apply").json()
    assert [x["id"] for x in res["affected_appointments"]] == [a["id"]]
    assert client.get(f"/appointments/{a['id']}", headers=admin).json()["appointment_time"] == "09:00"


def test_real_room_allocation_sample(client, admin):
    """The hospital's actual sample file: multi-line name cells (name / designation /
    'Mobile No. ...'), a per-row 'Floor Name' column ('OPD - (3rd Floor) Desk - A'), a
    Department column that isn't part of the schedule and is ignored, and room numbers
    that are only set on the first of a pair of rows sharing a room (vertical merge)."""
    for name in (
        "Prof. Dr. Md. Shahadat Hossain", "Dr. Mohammad Abdullah-Al-Amin",
        "Dr. Rukun Uddin Chowdhury", "Dr. Md. Azizur Rahman",
    ):
        client.post("/doctors", json={"name": name}, headers=admin)

    path = __import__("os").path.join(__import__("os").path.dirname(__file__), "fixtures", "RoomAllocation.xlsx")
    with open(path, "rb") as f:
        r = upload(client, admin, f.read(), "RoomAllocation.xlsx", "preview", default_max_patients=20)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"] == {
        "doctors_in_file": 4, "doctors_changed": 4, "sessions": 21,
        "errors": 0, "warnings": 0, "mobiles_removed": 4,
    }
    by = {d["doctor_name"]: d["sessions"] for d in body["doctors"]}
    # names must NOT include the designation or "Mobile No. ..." line
    assert set(by) == {
        "Prof. Dr. Md. Shahadat Hossain", "Dr. Mohammad Abdullah-Al-Amin",
        "Dr. Rukun Uddin Chowdhury", "Dr. Md. Azizur Rahman",
    }
    shahadat = {s["weekday"]: s for s in by["Prof. Dr. Md. Shahadat Hossain"]}
    assert set(shahadat) == {"saturday", "sunday", "monday", "wednesday", "thursday"}  # x on tue/fri
    assert (shahadat["saturday"]["start_time"], shahadat["saturday"]["end_time"]) == ("11:00", "14:00")
    assert (shahadat["saturday"]["room"], shahadat["saturday"]["floor"], shahadat["saturday"]["desk"]) == ("3044", "3rd", "A")
    # room number inherited via the vertical D2:D3 merge onto the second doctor's row
    amin = {s["weekday"]: s for s in by["Dr. Mohammad Abdullah-Al-Amin"]}
    assert amin["saturday"]["room"] == "3044"
    assert len(amin) == 6  # every day except friday
    # room inherited via the D4:D5 merge, even though this doctor's own Department differs
    azizur = {s["weekday"]: s for s in by["Dr. Md. Azizur Rahman"]}
    assert azizur["saturday"]["room"] == "3031"
    assert set(azizur) == {"saturday", "sunday", "monday", "tuesday", "wednesday"}  # x on thu/fri
    assert {p["code"] for p in body["problems"]} == {"MOBILE_REMOVED"}
