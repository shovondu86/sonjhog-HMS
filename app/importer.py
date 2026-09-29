"""OPD room-allocation import: parse a spreadsheet/CSV, validate, and (optionally)
replace each listed doctor's weekly schedule. Preview and apply share one code path;
preview simply rolls the transaction back."""
import csv
import io
import os
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, time, timedelta
from typing import Dict, List, Optional, Tuple

from sqlalchemy.orm import Session, joinedload

from app import audit
from app.enums import AppointmentStatus
from app.errors import ApiError
from app.models import Appointment, Doctor, DoctorSchedule, User
from app.scheduling import (
    Ctx, affected_for_doctor, overlap_time, overlap_validity, today, weekday_name,
)
from app.schemas import ScheduleOut

MAX_BYTES = 5 * 1024 * 1024
WEEKDAYS = ["saturday", "sunday", "monday", "tuesday", "wednesday", "thursday", "friday"]

_DAY_WORDS = {
    "saturday": "saturday", "sat": "saturday", "sunday": "sunday", "sun": "sunday",
    "monday": "monday", "mon": "monday", "tuesday": "tuesday", "tue": "tuesday",
    "tues": "tuesday", "wednesday": "wednesday", "wed": "wednesday",
    "thursday": "thursday", "thu": "thursday", "thur": "thursday", "thurs": "thursday",
    "friday": "friday", "fri": "friday",
}
_NO_SESSION = {"x", "-", "\u2013", "\u2014"}
_HONORIFICS = {"prof", "professor", "dr", "doctor"}

_HOURS = re.compile(
    r"^\s*(\d{1,2})(?:[:.](\d{2}))?\s*([ap])\.?m\.?\s*(?:to|-|\u2013|\u2014)\s*"
    r"(\d{1,2})(?:[:.](\d{2}))?\s*([ap])\.?m\.?\s*$",
    re.I,
)
_PHONE = re.compile(r"(?:\+?88[\s-]?)?0?1[3-9](?:[\s-]?\d){8}")
_PHONE_LABEL = re.compile(r"\b(?:mobile|mob|cell|phone|ph|tel|contact)\b\s*(?:no\.?)?\s*[:\-.]?", re.I)
_SECTION_FLOOR = re.compile(r"\(?\s*([\w]+)\s*floor\s*\)?", re.I)
_SECTION_DESK = re.compile(r"desk\s*[-:\u2013]?\s*([\w]+)", re.I)


@dataclass
class Cand:
    row: int
    doctor_name: Optional[str]
    doctor_id: Optional[int]
    weekday: str
    start: time
    end: time
    room: Optional[str] = None
    floor: Optional[str] = None
    desk: Optional[str] = None
    max_patients: Optional[int] = None
    slot_minutes: Optional[int] = None
    report_capacity: Optional[int] = None


@dataclass
class Problem:
    severity: str
    code: str
    detail: str
    row: Optional[int] = None
    doctor_name: Optional[str] = None
    weekday: Optional[str] = None

    def dump(self) -> dict:
        return {
            "severity": self.severity, "code": self.code, "row": self.row,
            "doctor_name": self.doctor_name, "weekday": self.weekday, "detail": self.detail,
        }


def unsupported(detail: str) -> ApiError:
    return ApiError(415, "UNSUPPORTED_FILE", detail)


# ── spreadsheet reading ──────────────────────────────────────────────────────
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, time):
        return v.strftime("%H:%M")
    return str(v).strip()


def read_grids(filename: str, content: bytes) -> List[Tuple[str, List[List[str]]]]:
    """Returns [(sheet_name, grid)] with merged cells expanded (the top-left value is
    copied into every cell of the merged range)."""
    ext = os.path.splitext((filename or "").lower())[1]
    try:
        if ext == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            out = []
            for ws in wb.worksheets:
                grid = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
                for rng in ws.merged_cells.ranges:
                    val = _cell(ws.cell(rng.min_row, rng.min_col).value)
                    for r in range(rng.min_row, rng.max_row + 1):
                        for c in range(rng.min_col, rng.max_col + 1):
                            while len(grid) < r:
                                grid.append([])
                            row = grid[r - 1]
                            while len(row) < c:
                                row.append("")
                            row[c - 1] = val
                out.append((ws.title, grid))
            return out
        if ext == ".xls":
            import xlrd

            book = xlrd.open_workbook(file_contents=content, formatting_info=True)
            out = []
            for sh in book.sheets():
                grid = [[_cell(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
                for rlo, rhi, clo, chi in sh.merged_cells:
                    val = grid[rlo][clo] if rlo < len(grid) and clo < len(grid[rlo]) else ""
                    for r in range(rlo, min(rhi, len(grid))):
                        for c in range(clo, chi):
                            while len(grid[r]) <= c:
                                grid[r].append("")
                            grid[r][c] = val
                out.append((sh.name, grid))
            return out
    except ApiError:
        raise
    except Exception as exc:  # corrupt / not really a spreadsheet
        raise unsupported(f"The file could not be read as a spreadsheet: {exc}")
    raise unsupported("Only .xls, .xlsx or .csv files are accepted.")


# ── spreadsheet parsing ──────────────────────────────────────────────────────
def _day_of(text: str) -> Optional[str]:
    return _DAY_WORDS.get(text.strip().lower().rstrip("."))


def _find_header(grid, start: int):
    for i in range(start, len(grid)):
        days = {c: _day_of(v) for c, v in enumerate(grid[i]) if v and _day_of(v)}
        if len(days) >= 3:
            return i, days
    return None, None


def _to_time(h: str, m: Optional[str], ap: str) -> time:
    hour = int(h) % 12 + (12 if ap.lower() == "p" else 0)
    return time(hour, int(m or 0))


def parse_hours(text: str):
    """-> list[(start, end)] | None when unparseable. 'A & B' is TWO sessions."""
    parts = [p for p in re.split(r"\s*&\s*|\s+and\s+|\s*;\s*|\s*,\s*|\n", text) if p.strip()]
    out = []
    for p in parts:
        m = _HOURS.match(p)
        if not m:
            return None
        start = _to_time(m.group(1), m.group(2), m.group(3))
        end = _to_time(m.group(4), m.group(5), m.group(6))
        if end <= start:
            return None
        out.append((start, end))
    return out or None


def _clean_name(raw: str) -> Tuple[str, int]:
    """The name cell is often multi-line: name, then designation, then a
    'Mobile No. 017...' line (e.g. "Dr. X\\nConsultant\\nMobile No. 01711525388").
    Only the first line is the name; designation/mobile lines are discarded entirely,
    never appended. Also drops any phone number that turns up inline on the name
    line itself, and 'Mobile:'-style labels."""
    lines = [ln.strip() for ln in re.split(r"[\r\n]+", raw) if ln.strip()]
    removed = sum(len(_PHONE.findall(ln)) for ln in lines)
    name = lines[0] if lines else ""
    name = _PHONE.sub(" ", name)
    name = _PHONE_LABEL.sub(" ", name)
    return re.sub(r"\s+", " ", name).strip(" ,;:-\u2013"), removed


def _first_line(raw: str) -> str:
    for ln in re.split(r"[\r\n]+", raw):
        if ln.strip():
            return ln.strip()
    return ""


def parse_grid(grid, sheet: str, multi: bool, cands: List[Cand], problems: List[Problem]) -> int:
    """Returns the number of phone numbers removed."""
    removed_total = 0
    hdr, day_cols = _find_header(grid, 0)
    if hdr is None:
        return 0
    first_day = min(day_cols)
    header = grid[hdr]
    # Prefer a column literally about the doctor before falling back to a generic "name" match
    # (a "Floor Name" or "Department" column can also contain the word "name").
    doc_col = next((c for c, v in enumerate(header) if "doctor" in v.lower()), None)
    if doc_col is None:
        doc_col = next(
            (c for c, v in enumerate(header)
             if any(k in v.lower() for k in ("name", "consultant", "physician"))),
            None,
        )
    room_col = next((c for c, v in enumerate(header) if "room" in v.lower()), None)
    # A per-row "Floor Name" column (e.g. "OPD - (3rd Floor) Desk - A") is the common case;
    # fall back to a standalone banner row (a row with only one filled cell) when there's no
    # such column, e.g. a sheet that puts the same text on its own row instead.
    floor_col = next((c for c, v in enumerate(header) if "floor" in v.lower()), None)
    if doc_col is None:
        doc_col = first_day - 1
    if room_col is None:
        room_col = doc_col - 1 if doc_col - 1 >= 0 else None
    where = f" (sheet '{sheet}')" if multi else ""
    banner_floor = banner_desk = None
    i = hdr + 1
    while i < len(grid):
        row = grid[i] + [""] * (max(day_cols) + 1 - len(grid[i]))
        rownum = i + 1
        i += 1
        # a repeated header row re-maps the columns
        days_here = {c: _day_of(v) for c, v in enumerate(row) if v and _day_of(v)}
        if len(days_here) >= 3:
            day_cols = days_here
            continue
        texts = {v for v in row if v}
        if not texts:
            continue
        # standalone banner row, e.g. "OPD - (3rd Floor) Desk - A" alone on its own row
        # (only used when there is no per-row Floor Name column)
        if floor_col is None and len(texts) == 1:
            t = next(iter(texts))
            f, d = _SECTION_FLOOR.search(t), _SECTION_DESK.search(t)
            if f or d:
                if f:
                    banner_floor = f.group(1)
                if d:
                    banner_desk = d.group(1)
                continue

        name_raw = row[doc_col] if doc_col is not None and doc_col < len(row) else ""
        room_raw = row[room_col] if room_col is not None and room_col < len(row) else ""
        room = _first_line(room_raw) if room_raw else ""
        name, removed = _clean_name(name_raw)
        if removed:
            removed_total += removed
            problems.append(Problem(
                "info", "MOBILE_REMOVED", f"A phone number in the name cell was dropped{where}.",
                row=rownum, doctor_name=name or None,
            ))
        if not name:
            if room:
                problems.append(Problem(
                    "warning", "FACILITY_ROW",
                    f"Room '{room}' has no doctor; skipped{where}.", row=rownum,
                ))
            continue

        if floor_col is not None:
            floor_text = row[floor_col] if floor_col < len(row) else ""
            f = _SECTION_FLOOR.search(floor_text) if floor_text else None
            d = _SECTION_DESK.search(floor_text) if floor_text else None
            floor = f.group(1) if f else None
            desk = d.group(1) if d else None
        else:
            floor, desk = banner_floor, banner_desk

        for col, wd in sorted(day_cols.items()):
            text = row[col] if col < len(row) else ""
            if text == "":
                problems.append(Problem(
                    "error", "BLANK_HOURS",
                    f"The {wd} cell is empty. Use x if there is no session{where}.",
                    row=rownum, doctor_name=name, weekday=wd,
                ))
                continue
            if text.strip().lower() in _NO_SESSION:
                continue
            spans = parse_hours(text)
            if spans is None:
                problems.append(Problem(
                    "error", "UNPARSEABLE_HOURS",
                    f"'{text}' is not in the form 'hh:mm am To hh:mm pm'{where}.",
                    row=rownum, doctor_name=name, weekday=wd,
                ))
                continue
            for s, e in spans:
                cands.append(Cand(
                    row=rownum, doctor_name=name, doctor_id=None, weekday=wd, start=s, end=e,
                    room=room or None, floor=floor, desk=desk,
                ))
    return removed_total


def parse_csv(content: bytes, cands: List[Cand], problems: List[Problem]) -> None:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise unsupported("The CSV must be UTF-8 encoded.")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise unsupported("The CSV has no header row.")
    reader.fieldnames = [f.strip().lower() for f in reader.fieldnames]
    need = {"weekday", "start_time", "end_time"}
    if not need <= set(reader.fieldnames) or not ({"doctor_id", "doctor_name"} & set(reader.fieldnames)):
        raise unsupported(
            "CSV columns: doctor_id (or doctor_name), weekday, start_time, end_time, room, "
            "floor, desk, max_patients, optional slot_minutes, report_capacity."
        )

    def val(r, k):
        v = (r.get(k) or "").strip()
        return v or None

    for n, r in enumerate(reader, start=2):
        if not any((v or "").strip() for v in r.values() if isinstance(v, str)):
            continue
        name = val(r, "doctor_name")
        did = val(r, "doctor_id")
        wd = _DAY_WORDS.get((val(r, "weekday") or "").lower())
        try:
            sh, sm = (val(r, "start_time") or "").split(":")[:2]
            eh, em = (val(r, "end_time") or "").split(":")[:2]
            start, end = time(int(sh), int(sm)), time(int(eh), int(em))
            if end <= start or wd is None:
                raise ValueError
        except ValueError:
            problems.append(Problem(
                "error", "UNPARSEABLE_HOURS",
                "weekday must be a day name; start_time/end_time 24-hour HH:MM with end after start.",
                row=n, doctor_name=name, weekday=wd,
            ))
            continue

        def as_int(k):
            v = val(r, k)
            return int(v) if v and v.isdigit() else None

        cands.append(Cand(
            row=n, doctor_name=name, doctor_id=int(did) if did and did.isdigit() else None,
            weekday=wd, start=start, end=end, room=val(r, "room"), floor=val(r, "floor"),
            desk=val(r, "desk"), max_patients=as_int("max_patients"),
            slot_minutes=as_int("slot_minutes"), report_capacity=as_int("report_capacity"),
        ))


# ── doctor resolution ────────────────────────────────────────────────────────
def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[.,]", " ", (s or "").lower())).strip()


def _core(s: str) -> str:
    return " ".join(w for w in _norm(s).split() if w not in _HONORIFICS)


def resolve_doctors(db: Session, cands: List[Cand], problems: List[Problem]) -> List[Cand]:
    docs = db.query(Doctor).filter(Doctor.deleted_at.is_(None)).all()
    by_id = {d.id: d for d in docs}
    keyed = [(d, _norm(d.name), _norm(d.name_bn or ""), _core(d.name), _core(d.name_bn or "")) for d in docs]
    cache: Dict[tuple, Optional[Doctor]] = {}
    reported = set()
    resolved = []
    for c in cands:
        key = (c.doctor_id, c.doctor_name)
        if key not in cache:
            doc = None
            if c.doctor_id is not None:
                doc = by_id.get(c.doctor_id)
                if doc is None:
                    problems.append(Problem(
                        "error", "UNKNOWN_DOCTOR", f"No doctor with id {c.doctor_id}.",
                        row=c.row, doctor_name=c.doctor_name,
                    ))
            else:
                n, k = _norm(c.doctor_name or ""), _core(c.doctor_name or "")
                exact = [d for d, a, b, _, _ in keyed if n and n in (a, b)]
                loose = [d for d, _, _, a, b in keyed if k and k in (a, b)]
                hits = exact or loose
                if len(hits) == 1:
                    doc = hits[0]
                elif len(hits) > 1:
                    problems.append(Problem(
                        "error", "AMBIGUOUS_DOCTOR",
                        f"'{c.doctor_name}' matches {len(hits)} doctors "
                        f"(ids {', '.join(str(h.id) for h in hits)}).",
                        row=c.row, doctor_name=c.doctor_name,
                    ))
                else:
                    problems.append(Problem(
                        "error", "UNKNOWN_DOCTOR", f"No doctor named '{c.doctor_name}'.",
                        row=c.row, doctor_name=c.doctor_name,
                    ))
            cache[key] = doc
        doc = cache[key]
        if doc is not None:
            c.doctor_id = doc.id
            c.doctor_name = doc.name
            resolved.append(c)
    return resolved


# ── main pipeline ────────────────────────────────────────────────────────────
def _tuple(weekday, start, end, room, floor, desk, max_p, slot, rep):
    return (weekday, start, end, room or None, floor or None, desk or None, max_p, slot or None, rep or 0)


def run_import(
    db: Session, user: User, *, filename: str, content: bytes, mode: str,
    valid_from: date, default_max_patients: Optional[int],
) -> Tuple[dict, List[Problem]]:
    cands: List[Cand] = []
    problems: List[Problem] = []
    mobiles = 0

    ext = os.path.splitext((filename or "").lower())[1]
    if ext == ".csv":
        parse_csv(content, cands, problems)
    else:
        grids = read_grids(filename, content)
        for name, grid in grids:
            mobiles += parse_grid(grid, name, len(grids) > 1, cands, problems)
        if not cands and not problems:
            raise unsupported(
                "No weekday header row (Saturday..Friday) was found in the spreadsheet."
            )

    doctors_in_file = len({(c.doctor_id, (c.doctor_name or "").lower()) for c in cands})
    cands = resolve_doctors(db, cands, problems)
    by_doc: Dict[int, List[Cand]] = defaultdict(list)
    for c in cands:
        by_doc[c.doctor_id].append(c)

    # current rows in effect from valid_from on, per doctor
    current: Dict[int, List[DoctorSchedule]] = defaultdict(list)
    for r in db.query(DoctorSchedule).filter(DoctorSchedule.is_active.is_(True)):
        if r.valid_to is None or r.valid_to >= valid_from:
            current[r.doctor_id].append(r)

    # capacity precedence: file -> doctor's current capacity for that weekday -> default
    sessions: Dict[int, list] = {}
    for did, cs in by_doc.items():
        out = []
        for c in cs:
            cur = sorted(
                (r for r in current.get(did, []) if r.weekday == c.weekday),
                key=lambda r: r.start_time,
            )
            cap = c.max_patients or (cur[0].max_patients if cur else None) or default_max_patients
            if not cap:
                problems.append(Problem(
                    "error", "MISSING_CAPACITY",
                    "No max_patients in the file, no current capacity for that weekday, and no "
                    "default_max_patients.",
                    row=c.row, doctor_name=c.doctor_name, weekday=c.weekday,
                ))
                continue
            slot = c.slot_minutes if c.slot_minutes is not None else (cur[0].slot_minutes if cur else None)
            rep = c.report_capacity if c.report_capacity is not None else (
                (cur[0].report_capacity if cur else 0) or 0
            )
            out.append((c, cap, slot, rep))
        sessions[did] = out

    # overlaps inside the file (same doctor) and room clashes (file vs file, file vs others)
    for did, items in sessions.items():
        for a in range(len(items)):
            for b in range(a + 1, len(items)):
                x, y = items[a][0], items[b][0]
                if x.weekday == y.weekday and overlap_time(x.start, x.end, y.start, y.end):
                    problems.append(Problem(
                        "error", "SCHEDULE_OVERLAP",
                        f"Two sessions overlap on {x.weekday} (rows {x.row} and {y.row}).",
                        row=y.row, doctor_name=y.doctor_name, weekday=y.weekday,
                    ))
    flat = [(did, it) for did, items in sessions.items() for it in items]
    for a in range(len(flat)):
        for b in range(a + 1, len(flat)):
            (d1, (x, *_)), (d2, (y, *_)) = flat[a], flat[b]
            if d1 != d2 and _room_clash(x, y):
                problems.append(Problem(
                    "error", "ROOM_CONFLICT",
                    f"Room {x.room} is given to two doctors on {x.weekday} at overlapping times "
                    f"(rows {x.row} and {y.row}).",
                    row=y.row, doctor_name=y.doctor_name, weekday=y.weekday,
                ))
    others = [
        r for did, rows in current.items() if did not in sessions for r in rows
    ]
    for did, items in sessions.items():
        for (c, *_r) in items:
            for r in others:
                if (
                    r.weekday == c.weekday and c.room and r.room
                    and c.room.lower() == r.room.lower()
                    and (c.desk or "").lower() == (r.desk or "").lower()
                    and overlap_time(c.start, c.end, r.start_time, r.end_time)
                    and overlap_validity(valid_from, None, r.valid_from, r.valid_to)
                ):
                    problems.append(Problem(
                        "error", "ROOM_CONFLICT",
                        f"Room {c.room} is already allocated to doctor {r.doctor_id} on {c.weekday}.",
                        row=c.row, doctor_name=c.doctor_name, weekday=c.weekday,
                    ))

    # active doctors with a schedule who are absent from the file are left untouched
    in_file = set(sessions)
    named = {d.id: d.name for d in db.query(Doctor).filter(Doctor.deleted_at.is_(None), Doctor.is_active.is_(True))}
    for did, rows in current.items():
        if did not in in_file and did in named and rows:
            problems.append(Problem(
                "warning", "NOT_IN_FILE",
                f"{named[did]} has a schedule but is not in the file; left as it is.",
                doctor_name=named[did],
            ))

    errors = [p for p in problems if p.severity == "error"]

    # Build the per-doctor result (changed?) ---------------------------------
    imported = []
    changed_ids = []
    plan: Dict[int, list] = {}
    for did, items in sessions.items():
        if not items:  # every session of this doctor failed validation (reported above)
            continue
        new_set = sorted(
            _tuple(c.weekday, c.start, c.end, c.room, c.floor, c.desk, cap, slot, rep)
            for c, cap, slot, rep in items
        )
        old_set = sorted(
            _tuple(r.weekday, r.start_time, r.end_time, r.room, r.floor, r.desk,
                   r.max_patients, r.slot_minutes, r.report_capacity)
            for r in current.get(did, [])
        )
        changed = new_set != old_set
        if changed:
            changed_ids.append(did)
        plan[did] = new_set
        imported.append({
            "doctor_id": did,
            "doctor_name": by_doc[did][0].doctor_name,
            "sessions": [
                {
                    "weekday": w, "start_time": s, "end_time": e, "room": rm, "floor": fl,
                    "desk": dk, "max_patients": cap, "slot_minutes": sl, "report_capacity": rp,
                    "valid_from": valid_from, "valid_to": None, "is_active": True,
                }
                for (w, s, e, rm, fl, dk, cap, sl, rp) in new_set
            ],
            "changed": changed,
        })

    result = {
        "mode": mode, "applied": False, "valid_from": valid_from, "doctors": imported,
        "problems": problems, "affected_appointments": [],
        "summary": {
            "doctors_in_file": doctors_in_file, "doctors_changed": len(changed_ids),
            "sessions": sum(len(v) for v in plan.values()),
            "errors": len(errors),
            "warnings": len([p for p in problems if p.severity == "warning"]),
            "mobiles_removed": mobiles,
        },
    }
    if errors:
        return result, problems

    # Do the work inside the transaction; preview rolls it back ---------------
    affected_rows: List[Appointment] = []
    for did in changed_ids:
        old = list(current.get(did, []))
        old_ids = {r.id for r in old}
        before = {"schedules": [_dump(r) for r in old]}
        for r in old:
            if r.valid_from is None or r.valid_from < valid_from:
                r.valid_to = valid_from - timedelta(days=1)
            else:  # a not-yet-started rota that this file supersedes
                r.is_active = False
        new_rows = []
        for (w, s, e, rm, fl, dk, cap, sl, rp) in plan[did]:
            row = DoctorSchedule(
                doctor_id=did, weekday=w, start_time=s, end_time=e, room=rm, floor=fl, desk=dk,
                max_patients=cap, slot_minutes=sl, report_capacity=rp,
                valid_from=valid_from, valid_to=None, is_active=True,
            )
            db.add(row)
            new_rows.append(row)
        db.flush()
        stale = _relink(db, did, old_ids, new_rows, valid_from)
        after = {"schedules": [_dump(r) for r in new_rows]}
        if mode == "apply":
            audit.record(db, user, "import", "doctor", did, before, after, doctor_id=did)
        affected_rows.extend(
            affected_for_doctor(db, did, max(valid_from, today()), unlinked_schedule_ids=stale)
        )

    result["affected_appointments"] = affected_rows
    result["applied"] = mode == "apply"
    return result, problems


def _room_clash(x: Cand, y: Cand) -> bool:
    return bool(
        x.weekday == y.weekday and x.room and y.room and x.room.lower() == y.room.lower()
        and (x.desk or "").lower() == (y.desk or "").lower()
        and overlap_time(x.start, x.end, y.start, y.end)
    )


def _dump(r: DoctorSchedule) -> dict:
    return ScheduleOut.model_validate(r).model_dump(mode="json")


def _relink(db: Session, doctor_id: int, old_ids: set, new_rows: list, valid_from: date) -> set:
    """Future bookings stay attached to a session: an appointment booked into an end-dated
    row is re-pointed at the new row that contains its time, when that mapping is
    unambiguous. Serial, time and date are never changed. Returns the ids of old rows whose
    bookings could NOT be re-linked, so they are reported as affected."""
    if not old_ids:
        return set()
    appts = (
        db.query(Appointment)
        .filter(
            Appointment.doctor_id == doctor_id,
            Appointment.schedule_id.in_(old_ids),
            Appointment.appointment_date >= valid_from,
            Appointment.status != AppointmentStatus.cancelled,
        )
        .all()
    )
    groups: Dict[tuple, list] = defaultdict(list)
    for a in appts:
        groups[(a.schedule_id, a.appointment_date, a.queue)].append(a)

    targets: Dict[tuple, int] = {}
    claimed: Dict[tuple, list] = defaultdict(list)
    for gkey, items in groups.items():
        wd = weekday_name(gkey[1])
        cands = [n for n in new_rows if n.weekday == wd]
        picks = set()
        for a in items:
            match = [n for n in cands if n.start_time <= a.appointment_time <= n.end_time]
            pref = [n for n in match if n.start_time == a.session_start] or match
            picks.add(pref[0].id if pref else None)
        if len(picks) == 1 and None not in picks:
            tid = next(iter(picks))
            targets[gkey] = tid
            claimed[(tid, gkey[1], gkey[2])].append(gkey)

    for (tid, d, q), gkeys in claimed.items():
        if len(gkeys) > 1:  # two old sessions collapse into one: serials would collide
            for g in gkeys:
                targets.pop(g, None)

    stale = set()
    for gkey, items in groups.items():
        if gkey in targets:
            for a in items:
                a.schedule_id = targets[gkey]
        else:
            stale.add(gkey[0])
    db.flush()
    return stale
